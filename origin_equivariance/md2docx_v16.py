# -*- coding: utf-8 -*-
"""
BSPC (Elsevier, Biomedical Signal Processing and Control) 单栏双倍行距 docx 生成器
==============================================================================
符合 BSPC Guide for Authors：
  - A4，边距 2.5 cm，单栏，Times New Roman 12 pt，双倍行距
  - 连续行号 + 页脚页码
  - 摘要 < 250 词，单段不分结构
  - 关键词 1–7 个
  - 章节阿拉伯数字编号（1. / 1.1.）
  - 图插正文、图注在图下；表注在表上，三线表（无竖线、无底纹）
  - 参考文献方括号数字，Elsevier 格式
  - CRediT / 利益冲突 / GenAI / 数据可用性 / 致谢
用法：python md2docx_bspc.py [--with-figs]
"""
import os, io, re, sys
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

D = r"G:\Work_work\math\EEG\paper_v8"
sys.path.insert(0, r"G:\Work_work\math\EEG\local_analysis")
from tnsre_render import iter_runs

SRC = os.path.join(D, "_v16_docx.md")
WITH_FIGS = "--with-figs" in sys.argv
OUT = os.path.join(D, "BSPC_Manuscript_v16_with_figs.docx" if WITH_FIGS
                   else "BSPC_Manuscript_v16.docx")
FIGD = os.path.join(D, "figs")

FONT = "Times New Roman"
TEXT_W = 6.30      # in，A4 减 2.5cm 边距
BODY = 12
TABLE_SZ = 9
CAP = 10.5

# ================================================================ 工具
def set_font(run, size=BODY, bold=False, italic=False, mono=False):
    name = "Courier New" if mono else FONT
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)


def add_line_numbers(section):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:lnNumType")):
        sectPr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    pgMar = sectPr.find(qn("w:pgMar"))
    if pgMar is not None:
        pgMar.addnext(ln)
    else:
        sectPr.insert(0, ln)


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r2 = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r2.append(t)
    fld.append(r2)
    run._element.addnext(fld)


def para(text="", size=BODY, align="l", bold=False, italic=False,
         before=0, after=0, indent=0.0, hang=0.0, line=2.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.first_line_indent = Inches(-hang if hang else indent)
    pf.left_indent = Inches(hang)
    pf.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "j": WD_ALIGN_PARAGRAPH.JUSTIFY,
                    "c": WD_ALIGN_PARAGRAPH.CENTER,
                    "r": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if text:
        add_runs(p, text, size=size, bold=bold, italic=italic)
    return p


def add_runs(p, text, size=BODY, bold=False, italic=False):
    for seg, fl in iter_runs(text):
        if not seg:
            continue
        r = p.add_run(seg)
        set_font(r, size, bold=(bold or fl["bold"]), italic=(italic or fl["italic"]),
                 mono=fl["mono"])
        if fl["sub"]:
            r.font.subscript = True
        if fl["sup"]:
            r.font.superscript = True


def heading(num, title, level=1):
    txt = ("%s. %s" % (num, title)) if num else title
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(18 if level == 1 else 12)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    r = p.add_run(txt)
    set_font(r, BODY, bold=True, italic=(level == 2))
    return p


def add_equation(body, num):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    tabs = pf.tab_stops
    tabs.add_tab_stop(Inches(TEXT_W / 2), WD_ALIGN_PARAGRAPH.CENTER)
    tabs.add_tab_stop(Inches(TEXT_W), WD_ALIGN_PARAGRAPH.RIGHT)
    r = p.add_run("\t")
    set_font(r, BODY)
    for seg, fl in iter_runs("$" + body + "$"):
        if not seg:
            continue
        r = p.add_run(seg)
        set_font(r, BODY, italic=True)
        if fl["sub"]:
            r.font.subscript = True
        if fl["sup"]:
            r.font.superscript = True
    r = p.add_run("\t(%s)" % num)
    set_font(r, BODY)


def set_three_line(t):
    tbl = t._tbl
    tblPr = tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge, val, sz in (("top", "single", "10"), ("bottom", "single", "10"),
                          ("left", "none", "0"), ("right", "none", "0"),
                          ("insideH", "none", "0"), ("insideV", "none", "0")):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000" if val == "single" else "auto")
        b.append(el)
    look = tblPr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(b)
    else:
        tblPr.append(b)


def cell_bottom_border(cell, sz="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    b = OxmlElement("w:tcBorders")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    b.append(el)
    tcPr.append(b)


TAB_W = {
    # v7: 1=六管线(3) 2=四性质(3) 3=六管线BA(9) 4=spread(4)
    #     5=三层恢复(9) 6=SEED划分(4) 7=collapse(8) 8=calibration(7)
    1: [1.45, 2.40, 2.45],
    2: [2.05, 2.20, 2.05],
    3: [1.25, 0.42, 0.60, 0.52, 0.52, 0.60, 0.52, 0.55, 0.52],
    4: [1.45, 1.55, 1.55, 1.70],
    5: [0.62, 0.30, 0.66, 0.62, 0.44, 0.62, 0.44, 0.62, 0.44],
    6: [1.70, 1.40, 1.70, 1.45],
    7: [0.85, 0.40, 1.50, 1.10, 0.40, 0.40, 0.40, 0.40],
    8: [0.90, 0.85, 0.55, 0.60, 0.60, 0.65, 0.90],
}
DEFAULT_W = [1.20, 0.70, 0.70, 0.70, 0.70, 0.70, 0.70]
_tabcount = [0]

def add_table(rows, cap_label, cap_text):
    cp = doc.add_paragraph()
    cp.paragraph_format.space_before = Pt(12)
    cp.paragraph_format.space_after = Pt(4)
    cp.paragraph_format.line_spacing = 1.0
    cp.paragraph_format.keep_with_next = True
    r = cp.add_run(cap_label + ". "); set_font(r, CAP, bold=True)
    add_runs(cp, cap_text, size=CAP, bold=False)

    ncol = max(len(x) for x in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            cell.text = ""
            pp = cell.paragraphs[0]
            pp.paragraph_format.space_before = Pt(2)
            pp.paragraph_format.space_after = Pt(2)
            pp.paragraph_format.line_spacing = 1.0
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 1 else WD_ALIGN_PARAGRAPH.CENTER
            add_runs(pp, txt, size=TABLE_SZ, bold=(ri == 0))
            if ri == 0:
                cell_bottom_border(cell)
        if ri == 0:
            trPr = t.rows[0]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:tblHeader"))
    _tabcount[0] += 1
    widths = TAB_W.get(_tabcount[0])
    if len(widths or []) != ncol:
        widths = [TEXT_W / float(ncol)] * ncol
    if widths and len(widths) == ncol:
        # 归一化到版心宽度，避免总宽偏窄或溢出
        tot = float(sum(widths))
        if tot > 0:
            widths = [w * TEXT_W / tot for w in widths]
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for ci, w in enumerate(widths):
            for r in t.rows:
                r.cells[ci].width = Inches(w)
    set_three_line(t)
    para("", size=6, line=1.0, after=10)


def add_figure(img, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(img, width=Inches(TEXT_W - 0.1))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cp.paragraph_format.space_after = Pt(12)
        cp.paragraph_format.line_spacing = 1.0
        add_runs(cp, caption, size=CAP)


# ================================================================ 文档骨架
doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(BODY)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 2.0

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Cm(2.5))
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)
add_line_numbers(sec)
add_page_number_footer(sec)

# ================================================================ 解析
src = io.open(SRC, encoding="utf-8-sig").read()

def grab(start_pat, end_pat=None):
    m = re.search(start_pat, src, re.S | re.M)
    if not m:
        return ""
    rest = src[m.end():]
    if end_pat:
        m2 = re.search(end_pat, rest, re.S | re.M)
        if m2:
            return rest[:m2.start()]
    return rest

title = re.search(r"^# (.+)$", src, re.M).group(1).strip()
# 注意：grab 内部用了 re.S，故此处不能用 ".+"（会跨行吞掉整个文件），改用 [^\n]*
front = grab(r"^#[^\n]*", r"\n---\n")
abstract = grab(r"^## Abstract\s*$", r"^\*\*Keywords")
keywords = grab(r"^\*\*Keywords:\*\*", r"^---")
body = "## 1. Introduction\n" + grab(r"^## 1\. Introduction\s*$", r"^## CRediT authorship")
back = "## CRediT authorship contribution statement\n" + grab(
    r"^## CRediT authorship contribution statement\s*$", r"^## References\s*$")
refs = grab(r"^## References\s*$", r"^## Figure captions\s*$")
figcap = grab(r"^## Figure captions\s*$")

# ---------------------------------------------------------------- 题名页
para(title, size=16, align="l", bold=True, before=0, after=18, line=1.15)

auth_line = None
affs, corr, orcid, emails, equal_note = [], None, None, None, None
for ln in front.split("\n"):
    s = ln.strip()
    if not s or s.startswith("#"):
        continue
    if s.startswith("^") and "^" in s[1:]:
        affs.append(s)
    elif s.startswith("\\*"):
        corr = s.lstrip("\\* ").strip()
    elif s.upper().startswith("ORCID"):
        orcid = s
    elif s.startswith("E-mail addresses") or s.startswith("E-mail address"):
        emails = s
    elif s.startswith("†") or s.startswith("\u2020") or s.startswith("\\dagger"):
        equal_note = s.lstrip("\\daggers† ").strip()
    elif auth_line is None:
        auth_line = s

if auth_line:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    for tok in re.split(r"(\^[^\^]+\^)", auth_line):
        if not tok:
            continue
        m = re.fullmatch(r"\^([^\^]+)\^", tok)
        if m:
            r = p.add_run(m.group(1)); set_font(r, BODY); r.font.superscript = True
        else:
            r = p.add_run(tok); set_font(r, BODY)

for a in affs:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    m = re.match(r"\^([a-z])\^\s*(.*)$", a)
    if m:
        r = p.add_run(m.group(1)); set_font(r, 11); r.font.superscript = True
        r = p.add_run(" " + m.group(2)); set_font(r, 11, italic=True)
    else:
        r = p.add_run(a); set_font(r, 11, italic=True)

if equal_note:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("\u2020 "); set_font(r, 11); r.font.superscript = True
    r = p.add_run(equal_note); set_font(r, 11)

if corr:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(corr); set_font(r, 11)

if orcid:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(orcid); set_font(r, 11)

if emails:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(emails); set_font(r, 11)

# ---------------------------------------------------------------- 摘要 / 关键词
if abstract.strip():
    heading(None, "Abstract")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 2.0
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, re.sub(r"\s+", " ", abstract).strip(), size=BODY)

if keywords.strip():
    heading(None, "Keywords")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 2.0
    add_runs(p, re.sub(r"\s+", " ", keywords).strip(), size=BODY)

# ================================================================ 正文
lines = body.split("\n")
i = 0
n = len(lines)
SEC_RE = re.compile(r"^## (\d+)\.\s+(.+)$")
SUB_RE = re.compile(r"^### (\d+\.\d+)\.?\s+(.+)$")
FIG_RE = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)$")
TABCAP_RE = re.compile(r"^\*\*(Table \d+)\.\*\*\s*(.*)$")
NOTE_RE = re.compile(r"^\*([^*].*?)\*$")

pending_tabcap = None

while i < n:
    raw = lines[i]
    s = raw.strip()

    if not s or s == "---":
        i += 1
        continue

    m = SEC_RE.match(s)
    if m:
        heading(m.group(1), m.group(2))
        i += 1
        continue

    # 无编号二级标题（例如 "## Figure captions"）
    m0 = re.match(r"^## (.+)$", s)
    if m0:
        heading(None, m0.group(1).strip())
        i += 1
        continue

    m = SUB_RE.match(s)
    if m:
        heading(m.group(1), m.group(2), level=2)
        i += 1
        continue

    # 公式块
    if s.startswith("$$"):
        body_eq = s.strip("$").strip()
        tag = "1"
        mt = re.search(r"\\tag\{([^}]*)\}", body_eq)
        if mt:
            tag = mt.group(1)
            body_eq = body_eq[:mt.start()] + body_eq[mt.end():]
        add_equation(body_eq.strip(), tag)
        i += 1
        continue

    # 图
    m = FIG_RE.match(s)
    if m:
        rel = m.group(1).replace("/", os.sep)
        fp = os.path.join(D, rel)
        cap = None
        j = i + 1
        while j < n and not lines[j].strip():
            j += 1
        if j < n and lines[j].strip().startswith("**Fig."):
            cap = lines[j].strip()
            i = j
        if WITH_FIGS and os.path.exists(fp):
            add_figure(fp, cap)
        elif cap:
            para(cap, size=CAP, after=8, line=1.0)
        i += 1
        continue

    # 表标题
    m = TABCAP_RE.match(s)
    if m:
        pending_tabcap = (m.group(1), m.group(2).strip())
        i += 1
        continue

    # 表格
    if s.startswith("|"):
        rows = []
        while i < n and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                rows.append(cells)
            i += 1
        lbl, txt = pending_tabcap if pending_tabcap else ("Table 1.", "")
        pending_tabcap = None
        add_table(rows, lbl, txt)
        continue

    # 表注（整行斜体）
    if NOTE_RE.match(s) and len(s) > 40:
        para(s, size=CAP, after=10, line=1.0, italic=False)
        i += 1
        continue

    # 段落（合并到空行为止）
    buf = [s]
    i += 1
    while i < n:
        t = lines[i].strip()
        if (not t or t == "---" or SEC_RE.match(t) or SUB_RE.match(t)
                or t.startswith("$$") or t.startswith("|")
                or FIG_RE.match(t) or TABCAP_RE.match(t)):
            break
        buf.append(t)
        i += 1
    text = " ".join(buf)
    hang = 0.0
    if re.match(r"^\d+\.\s", text):
        hang = 0.25
    para(text, size=BODY, hang=hang)

# ================================================================ 后置声明
bk = back
cur = None
buf = []
def flush():
    global buf, cur
    if not buf:
        return
    txt = " ".join(x.strip() for x in buf if x.strip())
    if txt:
        para(txt, size=BODY)
    buf = []

for ln in bk.split("\n"):
    s = ln.strip()
    if not s or s == "---":
        flush()
        continue
    if s.startswith("## "):
        flush()
        heading(None, s[3:].strip())
        continue
    buf.append(s)
flush()

# ================================================================ 参考文献
heading(None, "References")
for ln in refs.split("\n"):
    s = ln.strip()
    if not s or s == "---":
        continue
    para(s, size=11, hang=0.3, line=1.5, after=4)

# ================================================================ 图题
if figcap.strip():
    heading(None, "Figure captions")
    for ln in figcap.split("\n"):
        s = ln.strip()
        if not s or s == "---":
            continue
        para(s, size=CAP, after=8, line=1.0)

doc.save(OUT)
print("saved:", OUT, os.path.getsize(OUT) // 1024, "KB")
