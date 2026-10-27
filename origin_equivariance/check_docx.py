# -*- coding: utf-8 -*-
"""Quick structural check of the generated docx."""
import io
from docx import Document

P = r"G:\Work_work\math\EEG\paper_v8\BSPC_Manuscript_v8_with_figs.docx"
d = Document(P)
paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
print("paragraphs:", len(paras))
print("tables:", len(d.tables))
print("inline images:", len(d.inline_shapes))
figs = [p for p in paras if p.startswith("Fig.")]
print("figure captions:", len(figs))
for f in figs:
    print("  ", f[:110])
print("---- first 12 paragraphs ----")
for p in paras[:12]:
    print("  |", p[:120])
body = "\n".join(paras)
print("---- checks ----")
print("placeholder TODO/XXX:", body.count("TODO") + body.count("XXX"))
print("has References:", "References" in body)
print("refs count:", sum(1 for p in paras if p.startswith("[") and "]" in p[:5]))
