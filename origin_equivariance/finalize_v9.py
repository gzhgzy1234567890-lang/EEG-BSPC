# -*- coding: utf-8 -*-
"""Final one-word trim, then rebuild the docx and the highlights file."""
import io
import os
import re
import subprocess

D = r"G:\Work_work\math\EEG\paper_v8"
P = os.path.join(D, "BSPC_draft_v9.md")
s = io.open(P, encoding="utf-8").read()
s = s.replace("We separate three questions.", "Separating three questions:")
io.open(P, "w", encoding="utf-8").write(s)
ab = re.search(r"## Abstract\s*(.*?)\n\*\*Keywords", s, re.S).group(1)
print("abstract words:", len(re.findall(r"[A-Za-z][A-Za-z\-']*", ab)))

# rebuild manuscript docx
r = subprocess.run(["python", "build_v9b.py"], cwd=D, capture_output=True, text=True)
print((r.stdout or "").strip()[-200:])

# rebuild highlights for v9
from docx import Document
from docx.shared import Pt
block = re.search(r"## Highlights\s*\n((?:- .*\n)+)", s).group(1)
bullets = [b[2:].strip() for b in block.strip().split("\n") if b.startswith("- ")]
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
doc.add_heading("Highlights", level=1)
for b in bullets:
    doc.add_paragraph(b, style="List Bullet")
out = os.path.join(D, "BSPC_Highlights_v9.docx")
doc.save(out)
print("saved", out, "| bullets", len(bullets), "max", max(len(b) for b in bullets))
