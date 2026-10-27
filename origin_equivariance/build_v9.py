# -*- coding: utf-8 -*-
"""Build BSPC_Manuscript_v9_with_figs.docx from BSPC_draft_v9.md."""
import io
import os
import re
import subprocess

D = r"G:\Work_work\math\EEG\paper_v8"
src = io.open(os.path.join(D, "BSPC_draft_v9.md"), encoding="utf-8").read()

# docx-only copy: drop the Highlights block and the trailing captions section
s = re.sub(r"## Highlights\s*\n(?:- .*\n)+", "", src)
s = re.sub(r"## Figure captions\s*\n.*\Z", "", s, flags=re.S)
io.open(os.path.join(D, "_v9_docx.md"), "w", encoding="utf-8").write(s)

b = io.open(os.path.join(D, "md2docx_v8.py"), encoding="utf-8").read()
b = b.replace("_v8_docx.md", "_v9_docx.md")
b = b.replace("BSPC_Manuscript_v8_with_figs.docx", "BSPC_Manuscript_v9_with_figs.docx")
b = b.replace("BSPC_Manuscript_v8.docx", "BSPC_Manuscript_v9.docx")
io.open(os.path.join(D, "md2docx_v9.py"), "w", encoding="utf-8").write(b)
r = subprocess.run(["python", "md2docx_v9.py", "--with-figs"], cwd=D,
                   capture_output=True, text=True)
print((r.stdout or "").strip()[-120:], (r.stderr or "").strip()[-200:])

from docx import Document
d = Document(os.path.join(D, "BSPC_Manuscript_v9_with_figs.docx"))
ps = [p.text.strip() for p in d.paragraphs if p.text.strip()]
print("paragraphs", len(ps), "tables", len(d.tables), "images", len(d.inline_shapes))
for p in ps[:6]:
    print("|", p[:110])
body = "\n".join(ps)
print("carets:", body.count("^"), "| url:", "origin_equivariance" in body,
      "| refs:", sum(1 for p in ps if re.match(r"^\[\d+\]", p)))
