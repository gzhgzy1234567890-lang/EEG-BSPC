# -*- coding: utf-8 -*-
"""Final trim and rebuild of v13."""
import io
import os
import re
import subprocess

D = r"G:\Work_work\math\EEG\paper_v8"
P = os.path.join(D, "BSPC_draft_v9.md")
s = io.open(P, encoding="utf-8").read()
s = s.replace("We ask what the choice is, and what it can buy.",
              "We ask what the choice is and what it can buy.")
s = s.replace("and we report both rather than the better one.", "and we report both.")
io.open(P, "w", encoding="utf-8").write(s)
ab = re.search(r"## Abstract\s*(.*?)\n\*\*Keywords", s, re.S).group(1)
print("abstract:", len(re.findall(r"[A-Za-z][A-Za-z\-']*", ab)))

d = re.sub(r"## Highlights\s*\n(?:- .*\n)+", "", s)
d = re.sub(r"(^\*\*Keywords:\*\*.*$)", r"\1\n\n---", d, flags=re.M)
d = re.sub(r"## Figure captions\s*\n.*\Z", "", d, flags=re.S)
io.open(os.path.join(D, "_v13_docx.md"), "w", encoding="utf-8").write(d)
r = subprocess.run(["python", "md2docx_v13.py", "--with-figs"], cwd=D,
                   capture_output=True, text=True)
print((r.stdout or "").strip()[-70:])
from docx import Document
doc = Document(os.path.join(D, "BSPC_Manuscript_v13_with_figs.docx"))
ps = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
print("paras", len(ps), "tables", len(doc.tables), "images", len(doc.inline_shapes),
      "carets", "\n".join(ps).count("^"))
