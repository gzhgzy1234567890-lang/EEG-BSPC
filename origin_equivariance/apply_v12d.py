# -*- coding: utf-8 -*-
"""v12 minor items: SEED-IV point in Fig. 2, highlights, scope sentence."""
import io
import os
import re
import subprocess

D = r"G:\Work_work\math\EEG\paper_v8"
P = os.path.join(D, "BSPC_draft_v9.md")
s = io.open(P, encoding="utf-8").read()

# 1. scope sentence
s = s.replace(
    "There are 36 + 108 + 45 + 320 + 336 + 90 = 935 cells, from 76 subject units",
    "The five development datasets give 36 + 108 + 45 + 320 + 336 = 845 cells and the held-out dataset adds 90, for 935 in total; the 845 are the only cells used for the main analyses and the sweep (Table 3). They come from 76 subject units")
io.open(P, "w", encoding="utf-8").write(s)

# 2. Fig. 2 with the held-out point
fig = r"""
import os, sys
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
HERE = r"%s"
sys.path.insert(0, HERE)
import pv8_common as C

plt.rcParams.update({"font.size": 8, "axes.labelsize": 8, "axes.titlesize": 8,
                     "legend.fontsize": 6.5, "figure.dpi": 200, "savefig.bbox": "tight",
                     "axes.spines.top": False, "axes.spines.right": False})
z = np.load(os.path.join(C.OUT, "screen.npz"))
ds = np.array([str(x) for x in z["ds"]]); rep = np.array([str(x) for x in z["rep"]])
zw = z["rz"]; pen = z["pen"]
fig, ax = plt.subplots(1, 2, figsize=(6.4, 2.5))
for u in sorted(set(zip(ds, rep))):
    m = (ds == u[0]) & (rep == u[1])
    ax[0].scatter(np.median(zw[m]), np.median(pen[m]), s=np.sqrt(m.sum()) * 6 + 8,
                  color="#888888", edgecolor="k", lw=0.3)
ax[0].scatter(1.00, -0.070, s=140, marker="*", color="#c0392b", edgecolor="k",
              lw=0.4, label="SEED-IV (held out)", zorder=5)
ax[0].axhline(0, color="k", lw=0.5, ls=":")
ax[0].set_xlabel(r"source-only statistic $z_w$")
ax[0].set_ylabel("penalty of the one-sided pipeline")
ax[0].set_title("(a) unit medians and the held-out point", loc="left")
ax[0].legend(frameon=False, loc="lower left")
ax[1].scatter(np.log10(np.maximum(zw, 1e-2)), pen, s=3, color="#888888", alpha=0.5)
ax[1].scatter(np.log10(1.00), -0.070, s=90, marker="*", color="#c0392b",
              edgecolor="k", lw=0.4, zorder=5)
ax[1].axhline(0, color="k", lw=0.5, ls=":")
ax[1].set_xlabel(r"$\log_{10} z_w$")
ax[1].set_ylabel("penalty")
ax[1].set_title("(b) every cell", loc="left")
fig.tight_layout()
for ext in ("pdf", "png"):
    fig.savefig(os.path.join(C.OUT, "..", "figs", "fig2.%s" % ext))
print("fig2 redone")
""" % D
io.open(os.path.join(D, "pv8_fig2.py"), "w", encoding="utf-8").write(fig)
r = subprocess.run(["python", "pv8_fig2.py"], cwd=D, capture_output=True, text=True)
print((r.stdout or "").strip()[-60:], (r.stderr or "").strip()[-120:])

# 3. rebuild docx and highlights
b = io.open(os.path.join(D, "md2docx_v12.py"), encoding="utf-8").read()
d = re.sub(r"## Highlights\s*\n(?:- .*\n)+", "", s)
d = re.sub(r"(^\*\*Keywords:\*\*.*$)", r"\1\n\n---", d, flags=re.M)
d = re.sub(r"## Figure captions\s*\n.*\Z", "", d, flags=re.S)
io.open(os.path.join(D, "_v12_docx.md"), "w", encoding="utf-8").write(d)
r = subprocess.run(["python", "md2docx_v12.py", "--with-figs"], cwd=D, capture_output=True, text=True)
print((r.stdout or "").strip()[-70:])

block = re.search(r"## Highlights\s*\n((?:- .*\n)+)", s).group(1)
bullets = [x[2:].strip() for x in block.strip().split("\n") if x.startswith("- ")]
from docx import Document
from docx.shared import Pt
doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(12)
doc.add_heading("Highlights", level=1)
for x in bullets:
    doc.add_paragraph(x, style="List Bullet")
doc.save(os.path.join(D, "BSPC_Highlights_v12.docx"))
print("highlights", len(bullets), max(len(x) for x in bullets))
