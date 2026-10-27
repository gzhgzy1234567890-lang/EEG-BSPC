# -*- coding: utf-8 -*-
"""Trim the v16 abstract and rebuild the docx and the highlights file."""
import io
import os
import re
import subprocess

D = r"G:\Work_work\math\EEG\paper_v8"
P = os.path.join(D, "BSPC_draft_v16.md")
s = io.open(P, encoding="utf-8").read()

s = s.replace("The audit covers 845 dataset-representation-subject-pair cells from five datasets, plus 90 cells from SEED-IV used once, as a held-out test, after every choice had been fixed.",
              "The audit covers 845 cells from five datasets plus 90 from SEED-IV, used once as a held-out test.")
s = s.replace(" that family is larger than the sample mean, and no uniqueness is claimed.",
              " that family is larger than the sample mean.")
s = s.replace("and we report the parameterisation that fits better and the one the model implies rather than choosing between them.",
              "and we report both the parameterisation that fits better and the one the model implies.")
io.open(P, "w", encoding="utf-8").write(s)
ab = re.search(r"## Abstract\s*(.*?)\n\*\*Keywords", s, re.S).group(1)
print("abstract:", len(re.findall(r"[A-Za-z][A-Za-z\-]*", ab)))

d = re.sub(r"## Highlights\s*\n(?:- .*\n)+", "", s)
d = re.sub(r"(^\*\*Keywords:\*\*.*$)", r"\1\n\n---", d, flags=re.M)
d = re.sub(r"## Figure captions\s*\n.*\Z", "", d, flags=re.S)
io.open(os.path.join(D, "_v16_docx.md"), "w", encoding="utf-8").write(d)
r = subprocess.run(["python", "md2docx_v16.py", "--with-figs"], cwd=D,
                   capture_output=True, text=True)
print((r.stdout or "").strip()[-70:])

from docx import Document
from docx.shared import Pt
block = re.search(r"## Highlights\s*\n((?:- .*\n)+)", s).group(1)
bullets = [x[2:].strip() for x in block.strip().split("\n") if x.startswith("- ")]
doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(12)
doc.add_heading("Highlights", level=1)
for b in bullets:
    doc.add_paragraph(b, style="List Bullet")
doc.save(os.path.join(D, "BSPC_Highlights_v16.docx"))
print("highlights", len(bullets), "max chars", max(len(b) for b in bullets))
