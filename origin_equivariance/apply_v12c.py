# -*- coding: utf-8 -*-
"""v12b: trim the abstract to 250 words and rebuild the docx."""
import io
import re
import os
import subprocess

D = r"G:\Work_work\math\EEG\paper_v8"
P = os.path.join(D, "BSPC_draft_v9.md")
s = io.open(P, encoding="utf-8").read()

new_ab = """Removing a session mean is the first step in cross-session EEG decoding, and two operations circulate under that description: centring both sessions, or centring one. We separate three questions. Equivariance: a correction is invariant to the coordinate origin if it translates by a translation-equivariant estimate of the target mean; this family contains more than the sample mean, and we do not claim uniqueness. Score-difference equivalence: two corrections shift every class score by the same amount exactly when their references differ by a vector whose projection onto the centred classifier weights vanishes; for a finite sample this is sufficient, not necessary, for identical predictions. Performance: a signed threshold-displacement law predicts the per-cell change of balanced accuracy for mean-calibrated linear classifiers (r = 0.99, MAE 0.013 on the BCI data), and we report where it fails.

The main experiments use 845 session pairs from five datasets; 90 further pairs from SEED-IV, opened only after the statistic was fixed, are used for one held-out test. An origin sweep on the 845 leaves the equivariant pipelines unchanged by 0.000 and drives the one-sided one to exact chance. A direction-controlled sweep shows that a norm-based statistic is direction-blind; a projected source-only statistic orders the damage across the six datasets (rho = −0.94, exact p = 0.017) and separates a severe loss with AUC 0.84 on the development cells. It does not survive as an operational rule: the threshold flags every held-out cell (specificity 0.000). We present it as an exploratory screen for feature families, not as a detector of risky sessions.
"""
ab_old = re.search(r"## Abstract\s*(.*?)\n\*\*Keywords", s, re.S).group(0)
s = s.replace(ab_old, "## Abstract\n\n" + new_ab + "\n**Keywords")
io.open(P, "w", encoding="utf-8").write(s)
ab = re.search(r"## Abstract\s*(.*?)\n\*\*Keywords", s, re.S).group(1)
print("abstract:", len(re.findall(r"[A-Za-z][A-Za-z\-']*", ab)))

# build
d = re.sub(r"## Highlights\s*\n(?:- .*\n)+", "", s)
d = re.sub(r"(^\*\*Keywords:\*\*.*$)", r"\1\n\n---", d, flags=re.M)
d = re.sub(r"## Figure captions\s*\n.*\Z", "", d, flags=re.S)
io.open(os.path.join(D, "_v12_docx.md"), "w", encoding="utf-8").write(d)
b = io.open(os.path.join(D, "md2docx_v11.py"), encoding="utf-8").read()
b = b.replace("_v11_docx.md", "_v12_docx.md")
b = b.replace("BSPC_Manuscript_v11_with_figs.docx", "BSPC_Manuscript_v12_with_figs.docx")
b = b.replace("BSPC_Manuscript_v11.docx", "BSPC_Manuscript_v12.docx")
io.open(os.path.join(D, "md2docx_v12.py"), "w", encoding="utf-8").write(b)
r = subprocess.run(["python", "md2docx_v12.py", "--with-figs"], cwd=D, capture_output=True, text=True)
print((r.stdout or "").strip()[-80:])

from docx import Document
doc = Document(os.path.join(D, "BSPC_Manuscript_v12_with_figs.docx"))
ps = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
body = "\n".join(ps)
print("paras", len(ps), "tables", len(doc.tables), "images", len(doc.inline_shapes),
      "carets", body.count("^"))
print("935 still in text:", body.count("935"), "| 845:", body.count("845"))
