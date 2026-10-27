# -*- coding: utf-8 -*-
"""Build the separate Highlights file and rebuild the manuscript docx."""
import io
import os
import re
import subprocess

from docx import Document
from docx.shared import Pt

D = r"G:\Work_work\math\EEG\paper_v8"
s = io.open(os.path.join(D, "BSPC_draft_v8.md"), encoding="utf-8").read()
block = re.search(r"## Highlights\s*\n((?:- .*\n)+)", s).group(1)
bullets = [b[2:].strip() for b in block.strip().split("\n") if b.startswith("- ")]
print("highlights:", len(bullets), "max length:", max(len(b) for b in bullets))

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
doc.add_heading("Highlights", level=1)
for b in bullets:
    doc.add_paragraph(b, style="List Bullet")
out = os.path.join(D, "BSPC_Highlights_v8.docx")
doc.save(out)
print("saved", out)

for cmd in (["python", "prep_docx2.py"], ["python", "md2docx_v8.py", "--with-figs"]):
    r = subprocess.run(cmd, cwd=D, capture_output=True, text=True)
    print(cmd[0], cmd[-1], "->", (r.stdout or "").strip()[-90:], r.stderr.strip()[-120:])
