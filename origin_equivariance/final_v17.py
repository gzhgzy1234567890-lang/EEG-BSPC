# -*- coding: utf-8 -*-
"""Close the three minor items and rebuild the deliverables."""
import io
import os
import re
import subprocess
from docx import Document
from docx.shared import Pt

D = r"G:\Work_work\math\EEG\paper_v8"
P = os.path.join(D, "BSPC_draft_v16.md")
s = io.open(P, encoding="utf-8").read()

old = """- The equivariant family of mean corrections is stated; uniqueness is not claimed.
- Score-difference equivalence is characterised by the centred classifier weights.
- A direction-controlled sweep separates the origin's projection from its norm.
- A Gaussian threshold-displacement model explains the sign for LDA and logistic.
- The source-only risk statistic ranks feature families but does not screen sessions."""
new = """- The equivariant family of mean corrections is stated; uniqueness is not claimed.
- Score-difference equivalence is characterised by the centred classifier weights.
- Only the origin's projection on the weights damages the decoder, not its norm.
- A threshold-displacement model explains the sign for LDA and logistic.
- Three released framework components take the mean from the training fold."""
s = s.replace(old, new)

s = s.replace("These ranges describe the eleven units studied and are not validated thresholds;",
              "These are the ranges observed on the eleven development units, not thresholds;")

io.open(P, "w", encoding="utf-8").write(s)

d = re.sub(r"## Highlights\s*\n(?:- .*\n)+", "", s)
d = re.sub(r"(^\*\*Keywords:\*\*.*$)", r"\1\n\n---", d, flags=re.M)
d = re.sub(r"## Figure captions\s*\n.*\Z", "", d, flags=re.S)
io.open(os.path.join(D, "_v16_docx.md"), "w", encoding="utf-8").write(d)
r = subprocess.run(["python", "md2docx_v16.py", "--with-figs"], cwd=D,
                   capture_output=True, text=True)
print((r.stdout or "").strip()[-60:])

block = re.search(r"## Highlights\s*\n((?:- .*\n)+)", s).group(1)
bullets = [x[2:].strip() for x in block.strip().split("\n") if x.startswith("- ")]
doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(12)
doc.add_heading("Highlights", level=1)
for b in bullets:
    doc.add_paragraph(b, style="List Bullet")
doc.save(os.path.join(D, "BSPC_Highlights_v16.docx"))
print("highlights", len(bullets), "max", max(len(b) for b in bullets))

doc = Document(os.path.join(D, "BSPC_Manuscript_v16_with_figs.docx"))
ab = re.search(r"## Abstract\s*(.*?)\n\*\*Keywords", s, re.S).group(1)
print("tables", len(doc.tables), "images", len(doc.inline_shapes),
      "abstract", len(re.findall(r"[A-Za-z][A-Za-z\-]*", ab)))
