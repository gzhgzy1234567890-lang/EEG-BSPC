# -*- coding: utf-8 -*-
"""Add the Fig. 4 image link and build the v10 docx."""
import io
import os
import re
import subprocess

D = r"G:\Work_work\math\EEG\paper_v8"
P = os.path.join(D, "BSPC_draft_v9.md")
s = io.open(P, encoding="utf-8").read()
if "![Figure 4]" not in s:
    s = s.replace("**Fig. 4.** (a) Predicted change of Equation (4)",
                  "![Figure 4](figs/fig4.png)\n**Fig. 4.** (a) Predicted change of Equation (4)", 1)
    io.open(P, "w", encoding="utf-8").write(s)

d = re.sub(r"## Highlights\s*\n(?:- .*\n)+", "", s)
d = re.sub(r"(^\*\*Keywords:\*\*.*$)", r"\1\n\n---", d, flags=re.M)
d = re.sub(r"## Figure captions\s*\n.*\Z", "", d, flags=re.S)
io.open(os.path.join(D, "_v10_docx.md"), "w", encoding="utf-8").write(d)

b = io.open(os.path.join(D, "md2docx_v9.py"), encoding="utf-8").read()
b = b.replace("_v9_docx.md", "_v10_docx.md")
b = b.replace("BSPC_Manuscript_v9_with_figs.docx", "BSPC_Manuscript_v10_with_figs.docx")
b = b.replace("BSPC_Manuscript_v9.docx", "BSPC_Manuscript_v10.docx")
io.open(os.path.join(D, "md2docx_v10.py"), "w", encoding="utf-8").write(b)
r = subprocess.run(["python", "md2docx_v10.py", "--with-figs"], cwd=D,
                   capture_output=True, text=True)
print((r.stdout or "").strip()[-90:])

from docx import Document
doc = Document(os.path.join(D, "BSPC_Manuscript_v10_with_figs.docx"))
ps = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
body = "\n".join(ps)
print("paras", len(ps), "tables", len(doc.tables), "images", len(doc.inline_shapes))
print("carets:", body.count("^"), "| figs:", [p[:18] for p in ps if p.startswith("Fig.")])
