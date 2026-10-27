# -*- coding: utf-8 -*-
"""v15b: the three remaining inconsistencies, then rebuild."""
import io
import os
import re
import subprocess

D = r"G:\Work_work\math\EEG\paper_v8"
P = os.path.join(D, "BSPC_draft_v9.md")
s = io.open(P, encoding="utf-8").read()

# 1. rho vs s
s = s.replace("ρ = |s| is unsigned and cannot determine this sign; the sign is carried by t*.",
              "The two displacements are different quantities — ρ uses μ_B and measures the zero-versus-none shift, while s uses μ_A and measures the zero-versus-mean shift — and neither is signed, so neither alone can determine this sign; the sign is carried by t*.")

# 2. Table 5 caption
s = s.replace("**Table 5.** The screen at the level of dataset-by-representation units, medians within a unit. z_T is the origin distance ||μ_A||/||sd_A|| of the source session; ρ is the displacement of Equation (2); penalty = zero − none; gain = mean − none. Sorted by z_T.",
              "**Table 5.** The source-only statistic at the level of dataset-by-representation units, medians within a unit. z_T = ||μ_A||/||sd_A||; z_w is Equation (2) with the scale taken from the centred scores; penalty = zero − none; gain = mean − none. Sorted by z_w.")

# 3. the 80-97 claim
s = re.sub(r"80[–-]97 per cent[^.]*\.", "the settings listed in Table 7.", s)

io.open(P, "w", encoding="utf-8").write(s)

d = re.sub(r"## Highlights\s*\n(?:- .*\n)+", "", s)
d = re.sub(r"(^\*\*Keywords:\*\*.*$)", r"\1\n\n---", d, flags=re.M)
d = re.sub(r"## Figure captions\s*\n.*\Z", "", d, flags=re.S)
io.open(os.path.join(D, "_v13_docx.md"), "w", encoding="utf-8").write(d)
r = subprocess.run(["python", "md2docx_v13.py", "--with-figs"], cwd=D,
                   capture_output=True, text=True)
print((r.stdout or "").strip()[-70:])
from docx import Document
doc = Document(os.path.join(D, "BSPC_Manuscript_v13_with_figs.docx"))
ps = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
print("paras", len(ps), "tables", len(doc.tables), "images", len(doc.inline_shapes))
