# -*- coding: utf-8 -*-
import io
import os
import re
import subprocess

D = r"G:\Work_work\math\EEG\paper_v8"
P = os.path.join(D, "BSPC_draft_v16.md")
s = io.open(P, encoding="utf-8").read()
s = s.replace("Performance: the failure of the one-sided form is a threshold displacement, whose sign is set by the threshold error the training session already carries rather than by the size of the displacement.",
              "Performance: its failure is a threshold displacement whose sign is set by the threshold error the training session already carries.")
s = s.replace("ranks feature families within the development data (AUC 0.84) but its cross-dataset association is rho = −0.77 (exact p = 0.10)",
              "ranks feature families (AUC 0.84) but its cross-dataset association is rho = −0.77 (p = 0.10)")
io.open(P, "w", encoding="utf-8").write(s)
ab = re.search(r"## Abstract\s*(.*?)\n\*\*Keywords", s, re.S).group(1)
print("abstract:", len(re.findall(r"[A-Za-z][A-Za-z\-]*", ab)))

d = re.sub(r"## Highlights\s*\n(?:- .*\n)+", "", s)
d = re.sub(r"(^\*\*Keywords:\*\*.*$)", r"\1\n\n---", d, flags=re.M)
d = re.sub(r"## Figure captions\s*\n.*\Z", "", d, flags=re.S)
io.open(os.path.join(D, "_v16_docx.md"), "w", encoding="utf-8").write(d)
r = subprocess.run(["python", "md2docx_v16.py", "--with-figs"], cwd=D,
                   capture_output=True, text=True)
print((r.stdout or "").strip()[-70:])
from docx import Document
doc = Document(os.path.join(D, "BSPC_Manuscript_v16_with_figs.docx"))
print("tables", len(doc.tables), "images", len(doc.inline_shapes),
      "paras", len([p for p in doc.paragraphs if p.text.strip()]))
