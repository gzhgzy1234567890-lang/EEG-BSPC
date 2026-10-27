# -*- coding: utf-8 -*-
"""Fix the v9 docx build: keywords separator, Fig. 3 placement, stray caret."""
import io
import os
import re
import subprocess

D = r"G:\Work_work\math\EEG\paper_v8"
P = os.path.join(D, "BSPC_draft_v9.md")
s = io.open(P, encoding="utf-8").read()

# 1. no literal caret in equation (5)
s = s.replace("m_k^A", "m_kA")

# 2. inline Figure 3 right after the per-subject gain sentence
if "![Figure 3]" not in s:
    cap3 = ("**Fig. 3.** (a) Balanced accuracy of the three pipelines by dataset; "
            "(b) the gain of the two-sided over the uncorrected pipeline, one point "
            "per subject, with the box showing the median and quartiles.")
    anchor = "Fig. 3 shows the levels by dataset and the gain per subject."
    s = s.replace(anchor, anchor + "\n\n![Figure 3](figs/fig3.png)\n" + cap3, 1)
    # remove the duplicate caption at the end
    s = re.sub(r"\n\*\*Fig\. 3\.\*\* The same.*?\n", "\n", s)
    s = re.sub(r"\n\*\*Fig\. 3\.\*\* \(a\) Balanced accuracy of the three pipelines by dataset; \(b\) the gain of the two-sided over the uncorrected pipeline, one point per subject, with the box showing the median and quartiles\.\n", "\n", s)

io.open(P, "w", encoding="utf-8").write(s)

# 3. docx-only copy with a '---' separator after the keywords line
src = io.open(P, encoding="utf-8").read()
d = re.sub(r"## Highlights\s*\n(?:- .*\n)+", "", src)
d = re.sub(r"(^\*\*Keywords:\*\*.*$)", r"\1\n\n---", d, flags=re.M)
d = re.sub(r"## Figure captions\s*\n.*\Z", "", d, flags=re.S)
io.open(os.path.join(D, "_v9_docx.md"), "w", encoding="utf-8").write(d)

r = subprocess.run(["python", "md2docx_v9.py", "--with-figs"], cwd=D,
                   capture_output=True, text=True)
print((r.stdout or "").strip()[-90:], (r.stderr or "").strip()[-150:])

from docx import Document
doc = Document(os.path.join(D, "BSPC_Manuscript_v9_with_figs.docx"))
ps = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
print("paras", len(ps), "tables", len(doc.tables), "images", len(doc.inline_shapes))
body = "\n".join(ps)
print("carets:", body.count("^"), "| figs:", [p[:40] for p in ps if p.startswith("Fig.")])
print("| keywords line:", [p[:80] for p in ps if p.startswith("cross-session EEG decoding")])
print("| intro present:", any(p.startswith("1. Introduction") for p in ps))
